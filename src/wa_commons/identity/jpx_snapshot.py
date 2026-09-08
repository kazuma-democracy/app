from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Iterable

from .jpx import from_jpx_row
from .models import EntityRecord, SourceRef
from .snapshots import sha256_file

DOMESTIC_MARKET_MARKER = "内国株式"
TSE_UNIVERSE_NORMALIZATION_VERSION = "tse-universe-normalization-v0.1"
IN_SCOPE_MARKETS = {
    "プライム（内国株式）": "Prime",
    "スタンダード（内国株式）": "Standard",
    "グロース（内国株式）": "Growth",
}


def _read_csv(path: Path) -> list[dict[str, object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def _rows_from_matrix(values: list[list[object]]) -> list[dict[str, object]]:
    if not values:
        return []
    headers = [str(v).strip() if v is not None else "" for v in values[0]]
    return [
        {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
        for row in values[1:]
    ]


def _read_xlsx(path: Path) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required to read .xlsx files") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    return _rows_from_matrix([list(row) for row in ws.iter_rows(values_only=True)])


def _read_xls(path: Path) -> list[dict[str, object]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("xlrd is required to read .xls files") from exc
    book = xlrd.open_workbook(path)
    sheet = book.sheet_by_index(0)
    return _rows_from_matrix([sheet.row_values(i) for i in range(sheet.nrows)])


def read_jpx_rows(path: str | Path) -> list[dict[str, object]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix == ".xls":
        return _read_xls(path)
    raise ValueError(f"unsupported JPX snapshot format: {suffix}")


def domestic_company_rows(rows: Iterable[dict[str, object]]) -> list[dict[str, object]]:
    out = []
    for row in rows:
        market = str(row.get("市場・商品区分", ""))
        if DOMESTIC_MARKET_MARKER not in market:
            continue
        code = str(row.get("コード", "")).strip()
        name = str(row.get("銘柄名", "")).strip()
        if not code or not name:
            continue
        out.append(row)
    return out


def _exclusion_category(market: str) -> str:
    if "ETF" in market or "ETN" in market:
        return "etf_etn"
    if "REIT" in market or "不動産投資信託" in market:
        return "reit"
    if "PRO Market" in market or "TOKYO PRO Market" in market:
        return "tokyo_pro_market"
    return "other_market"


def _normalize_snapshot_date(value: object) -> str:
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text.replace("-", "").replace("/", "")


def _semantic_projection(entities: list[dict[str, object]]) -> list[dict[str, str]]:
    return [
        {
            "entity_id": str(row["entity_id"]),
            "security_code": str(row["security_code"]),
            "canonical_name": str(row["canonical_name"]),
            "market_segment": str(row["market_segment"]),
        }
        for row in entities
    ]


def _semantic_payload_sha256(entities: list[dict[str, object]]) -> str:
    encoded = json.dumps(
        _semantic_projection(entities),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def build_universe(
    path: str | Path,
    *,
    snapshot: str,
    source_url: str,
    retrieved_at: str,
) -> dict:
    """Build the canonical TSE company universe from a locally supplied JPX snapshot.

    Row-level output is intended for the operator's authorized local workspace only.
    The manifest is deliberately non-row-level so it can be used as public
    verification metadata without redistributing the JPX-derived company list.
    """

    path = Path(path)
    expected_source_date = _normalize_snapshot_date(snapshot)
    source = SourceRef(
        source="JPX",
        source_key=path.name,
        snapshot=snapshot,
        url=source_url,
        retrieved_at=retrieved_at,
        adapter_version="0.4",
    )
    market_counts = {"Prime": 0, "Standard": 0, "Growth": 0}
    exclusion_counts = {
        "etf_etn": 0,
        "reit": 0,
        "tokyo_pro_market": 0,
        "other_market": 0,
    }
    entities: list[dict[str, object]] = []

    for row in read_jpx_rows(path):
        source_date = _normalize_snapshot_date(row.get("日付", ""))
        if not source_date:
            raise ValueError("source row date missing")
        if source_date != expected_source_date:
            raise ValueError(
                f"source row date {source_date} does not match snapshot {expected_source_date}"
            )

        market = str(row.get("市場・商品区分", "")).strip()
        segment = IN_SCOPE_MARKETS.get(market)
        if segment is None:
            exclusion_counts[_exclusion_category(market)] += 1
            continue

        code = str(row.get("コード", "")).strip()
        name = str(row.get("銘柄名", "")).strip()
        if not code or not name:
            raise ValueError("in-scope JPX row missing security code or name")

        entity = from_jpx_row(row, source)
        security_code = entity.identifiers[0].value
        entity_payload = entity.to_dict()
        entity_payload["security_code"] = security_code
        entity_payload["market_segment"] = segment
        entities.append(entity_payload)
        market_counts[segment] += 1

    entities.sort(key=lambda row: str(row["security_code"]))
    for previous, current in zip(entities, entities[1:]):
        if previous["security_code"] == current["security_code"]:
            raise ValueError(f"duplicate security_code: {current['security_code']}")

    return {
        "manifest": {
            "source": "JPX",
            "snapshot": snapshot,
            "source_url": source_url,
            "retrieved_at": retrieved_at,
            "source_file": path.name,
            "source_sha256": sha256_file(path),
            "adapter_version": "0.4",
            "normalization_rule_version": TSE_UNIVERSE_NORMALIZATION_VERSION,
            "selection": "TSE Prime/Standard/Growth domestic companies only",
            "rights_mode": "local_generation_only",
            "row_level_publication": "not_authorized_by_default_free_site_route",
            "market_counts": market_counts,
            "entity_count": len(entities),
            "exclusion_counts": exclusion_counts,
            "semantic_payload_sha256": _semantic_payload_sha256(entities),
        },
        "entities": entities,
    }


def write_universe(
    payload: dict,
    local_output: str | Path,
    public_manifest: str | Path,
) -> None:
    local_path = Path(local_output)
    manifest_path = Path(public_manifest)
    if local_path.resolve() == manifest_path.resolve():
        raise ValueError("local and public outputs must be distinct")
    local_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    local_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    manifest_path.write_text(
        json.dumps(payload["manifest"], ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def build_pilot(
    path: str | Path,
    *,
    snapshot: str,
    source_url: str,
    retrieved_at: str,
    limit: int = 100,
) -> dict:
    path = Path(path)
    rows = domestic_company_rows(read_jpx_rows(path))
    rows = sorted(rows, key=lambda r: str(r.get("コード", "")))[:limit]
    source = SourceRef(
        source="JPX",
        source_key=path.name,
        snapshot=snapshot,
        url=source_url,
        retrieved_at=retrieved_at,
        adapter_version="0.3",
    )
    entities: list[EntityRecord] = [from_jpx_row(row, source) for row in rows]
    return {
        "manifest": {
            "source": "JPX",
            "snapshot": snapshot,
            "source_url": source_url,
            "retrieved_at": retrieved_at,
            "source_file": path.name,
            "source_sha256": sha256_file(path),
            "adapter_version": "0.3",
            "selection": "sorted domestic listed equities by security code",
            "limit": limit,
            "entity_count": len(entities),
        },
        "entities": [entity.to_dict() for entity in entities],
    }


def write_pilot(payload: dict, output: str | Path) -> None:
    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
